"""Word (.docx) exports of saved schemes of work and daily lesson plans.

Built from exactly what is stored in Postgres -- nothing is generated or filled in here, so an
empty field stays empty in the document. Written against the real stored shapes
(`schemes_of_work.content.rows` = TermPlanRow[], `lesson_plans.content` = LessonPlanDraft), not the
old CLI models in models.py / doc_generator.py.
"""
import io
import re
from datetime import date

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

MUTED = RGBColor(0x55, 0x55, 0x55)

# Column order is fixed by the spec. Widths (inches) fit landscape A4.
SCHEME_COLUMNS = [
    ("Week", "week", 0.5),
    ("Lessons", "lessons", 0.6),
    ("Strand", "strand", 0.9),
    ("Sub-strand", "subStrand", 1.0),
    ("Key Inquiry Question", "keyInquiryQuestion", 1.1),
    ("Specific Learning Outcomes", "outcomes", 1.6),
    ("Suggested Learning Experiences", "experiences", 1.7),
    ("Resources", "resources", 1.05),
    ("Assessment", "assessment", 1.05),
    ("Reflection", "reflection", 0.95),
]


def safe_filename(*parts) -> str:
    cleaned = (re.sub(r"[^A-Za-z0-9]+", "_", str(part)).strip("_") for part in parts if part)
    return f"{'_'.join(part for part in cleaned if part) or 'document'}.docx"


def _text(value) -> str:
    return str(value or "").strip()


def _new_document(font_size: int, landscape: bool):
    doc = docx.Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(font_size)
    section = doc.sections[0]
    # A4, the paper size Kenyan schools print on.
    width, height = Inches(8.27), Inches(11.69)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        width, height = height, width
    section.page_width, section.page_height = width, height
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.5 if landscape else 0.8))
    return doc


def _add_title(doc, title: str, confirmed: bool):
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run("Confirmed teacher work product. Not official KICD content." if confirmed
                         else "DRAFT: not yet confirmed by the teacher. Not official KICD content.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def _repeat_as_header(row):
    """Repeat this row at the top of every page the table spans."""
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(flag)


def _set_cell(cell, text: str, bold: bool = False, width=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    if width is not None:
        cell.width = width


def _to_bytes(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_scheme_docx(scheme_row: dict) -> bytes:
    """A `schemes_of_work` row -> landscape scheme-of-work table, one table row per TermPlanRow."""
    doc = _new_document(font_size=9, landscape=True)
    title = (f"Scheme of Work — {_text(scheme_row.get('subject'))}, {_text(scheme_row.get('grade'))}, "
             f"Term {_text(scheme_row.get('term'))} {_text(scheme_row.get('year'))}")
    _add_title(doc, title, scheme_row.get("status") == "confirmed")

    table = doc.add_table(rows=1, cols=len(SCHEME_COLUMNS))
    table.style = "Table Grid"
    table.autofit = False
    _repeat_as_header(table.rows[0])
    for cell, (heading, _field, width) in zip(table.rows[0].cells, SCHEME_COLUMNS):
        _set_cell(cell, heading, bold=True, width=Inches(width))

    for row in (scheme_row.get("content") or {}).get("rows") or []:
        cells = table.add_row().cells
        for cell, (_heading, field, width) in zip(cells, SCHEME_COLUMNS):
            _set_cell(cell, _text(row.get(field)), width=Inches(width))

    return _to_bytes(doc)


def _format_date(value) -> str:
    if isinstance(value, date):
        return value.strftime("%A, %d %B %Y")
    try:
        return date.fromisoformat(_text(value)).strftime("%A, %d %B %Y")
    except ValueError:
        return _text(value)


def _add_section(doc, heading: str, text: str):
    doc.add_heading(heading, level=2)
    for line in _text(text).splitlines() or [""]:
        doc.add_paragraph(line)


def generate_lesson_docx(lesson_row: dict, scheme_row: dict | None) -> bytes:
    """A `lesson_plans` row (+ its parent scheme, if linked) -> daily lesson plan document."""
    lesson = lesson_row.get("content") or {}
    doc = _new_document(font_size=11, landscape=False)
    _add_title(doc, "Daily Lesson Plan", lesson_row.get("status") == "confirmed")

    details = [
        ("Lesson title", _text(lesson.get("title"))),
        ("Date", _format_date(lesson.get("date") or lesson_row.get("lesson_date"))),
        ("Time allocation", _text(lesson.get("duration"))),
        ("Roll", _text(lesson.get("roll"))),
        ("Strand", _text(lesson_row.get("strand"))),
        ("Sub-strand", _text(lesson_row.get("sub_strand"))),
    ]
    # Teaching context only exists on a linked scheme; an unlinked lesson simply has no such lines.
    if scheme_row:
        details += [
            ("Grade", _text(scheme_row.get("grade"))),
            ("Subject", _text(scheme_row.get("subject"))),
            ("Term", f"Term {_text(scheme_row.get('term'))} {_text(scheme_row.get('year'))}"),
        ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in details:
        cells = table.add_row().cells
        _set_cell(cells[0], label, bold=True, width=Inches(1.8))
        _set_cell(cells[1], value, width=Inches(4.87))

    _add_section(doc, "Specific Learning Outcomes", lesson.get("outcomes"))
    _add_section(doc, "Key Inquiry Question", lesson.get("keyInquiryQuestion"))
    # Both teacher-entered in the lesson form; never generated.
    _add_section(doc, "Core Competencies", lesson.get("competencies"))
    _add_section(doc, "Values and PCIs", lesson.get("valuesAndPcis"))
    _add_section(doc, "Learning Resources", lesson.get("resources"))
    _add_section(doc, "Introduction", lesson.get("introduction"))

    doc.add_heading("Main Learning Activities", level=2)
    steps = [_text(step) for step in lesson.get("development") or [] if _text(step)]
    for step in steps:
        doc.add_paragraph(step, style="List Number")
    if not steps:
        doc.add_paragraph("")

    _add_section(doc, "Assessment Activity", lesson.get("assessmentActivity"))
    _add_section(doc, "Lesson Closure", lesson.get("conclusion"))
    _add_section(doc, "Teacher Notes", lesson.get("teacherNotes"))

    # LessonPlanDraft has no reflection field: it is written by hand after the lesson is taught.
    doc.add_heading("Reflection", level=2)
    for _ in range(4):
        doc.add_paragraph("_" * 90)

    return _to_bytes(doc)
